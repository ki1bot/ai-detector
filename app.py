import io
import re
from pathlib import Path
from typing import Optional

import joblib
import numpy as np
from docx import Document
from fastapi import (
    FastAPI,
    File,
    Form,
    HTTPException,
    Request,
    UploadFile,
)
from fastapi.responses import (
    HTMLResponse,
    Response,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pypdf import PdfReader

from features import (
    indonesian_signal,
    meta_score_features,
    split_sentences,
    style_snapshot,
    stylometry_matrix,
    tokenize_words,
)

BASE_DIR = (
    Path(__file__)
    .resolve()
    .parent
)

MODEL_PATH = (
    BASE_DIR
    / "models"
    / "detector.joblib"
)

MAX_FILE_BYTES = (
    8
    * 1024
    * 1024
)

ALLOWED_EXTENSIONS = {
    ".txt",
    ".md",
    ".docx",
    ".pdf",
}

FORMAT_VERSION = 2

app = FastAPI(
    title="AI Detector Indonesia",
    version="2.0.0",
)

app.mount(
    "/static",
    StaticFiles(
        directory=(
            BASE_DIR
            / "static"
        )
    ),
    name="static",
)

templates = Jinja2Templates(
    directory=(
        BASE_DIR
        / "templates"
    )
)

_bundle = None


def get_bundle():
    global _bundle

    if _bundle is None:
        if not MODEL_PATH.exists():
            raise HTTPException(
                status_code=503,
                detail=(
                    "Model belum tersedia. "
                    "Jalankan python train.py terlebih dahulu."
                ),
            )

        loaded = joblib.load(
            MODEL_PATH
        )

        if (
            loaded.get(
                "format_version"
            )
            != FORMAT_VERSION
        ):
            raise HTTPException(
                status_code=503,
                detail=(
                    "Model lama tidak kompatibel. "
                    "Hapus models/detector.joblib "
                    "lalu jalankan python train.py lagi."
                ),
            )

        _bundle = loaded

    return _bundle


def normalize_text(
    text: str,
) -> str:
    text = (
        str(text)
        .replace(
            "\u00a0",
            " ",
        )
        .replace(
            "\u200b",
            " ",
        )
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def count_words(
    text: str,
) -> int:
    return len(
        tokenize_words(
            text
        )
    )


def extract_docx(
    content: bytes,
) -> str:
    document = Document(
        io.BytesIO(
            content
        )
    )

    blocks = []

    for paragraph in document.paragraphs:
        value = (
            paragraph.text
            .strip()
        )

        if value:
            blocks.append(
                value
            )

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                blocks.append(
                    " | ".join(
                        values
                    )
                )

    return "\n\n".join(
        blocks
    )


def extract_pdf(
    content: bytes,
) -> str:
    reader = PdfReader(
        io.BytesIO(
            content
        )
    )

    pages = []

    for page in reader.pages:
        value = (
            page.extract_text()
            or ""
        )

        value = re.sub(
            r"[ \t]+",
            " ",
            value,
        ).strip()

        if value:
            pages.append(
                value
            )

    return "\n\n".join(
        pages
    )


def extract_text(
    filename: str,
    content: bytes,
) -> str:
    suffix = (
        Path(filename)
        .suffix
        .lower()
    )

    if (
        suffix
        not in ALLOWED_EXTENSIONS
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Format file harus "
                "TXT, MD, DOCX, atau PDF."
            ),
        )

    try:
        if suffix in {
            ".txt",
            ".md",
        }:
            return content.decode(
                "utf-8",
                errors="replace",
            )

        if suffix == ".docx":
            return extract_docx(
                content
            )

        return extract_pdf(
            content
        )

    except HTTPException:
        raise

    except Exception as exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "Dokumen tidak dapat dibaca: "
                f"{exception}"
            ),
        ) from exception


def chunk_text(
    text: str,
    target_words: int = 135,
    max_words: int = 190,
) -> list[str]:
    paragraphs = [
        paragraph.strip()
        for paragraph in re.split(
            r"\n\s*\n",
            text,
        )
        if paragraph.strip()
    ]

    if not paragraphs:
        paragraphs = [
            text
        ]

    chunks = []
    current = []
    current_words = 0

    def flush():
        nonlocal current
        nonlocal current_words

        if current:
            value = " ".join(
                current
            ).strip()

            if value:
                chunks.append(
                    value
                )

        current = []
        current_words = 0

    for paragraph in paragraphs:
        sentences = (
            split_sentences(
                paragraph
            )
            or [paragraph]
        )

        for sentence in sentences:
            words = tokenize_words(
                sentence
            )

            if not words:
                continue

            if (
                len(words)
                > max_words
            ):
                flush()

                raw_words = (
                    sentence.split()
                )

                for start in range(
                    0,
                    len(raw_words),
                    target_words,
                ):
                    piece = " ".join(
                        raw_words[
                            start:
                            start + target_words
                        ]
                    ).strip()

                    if piece:
                        chunks.append(
                            piece
                        )

                continue

            if (
                current
                and current_words
                + len(words)
                > max_words
            ):
                flush()

            current.append(
                sentence
            )

            current_words += len(
                words
            )

            if (
                current_words
                >= target_words
            ):
                flush()

    flush()

    return (
        chunks
        or [text]
    )


def predict_texts(
    texts: list[str],
    bundle: dict,
) -> tuple[
    np.ndarray,
    np.ndarray,
]:
    models = bundle[
        "models"
    ]

    word_scores = (
        models["word"]
        .predict_proba(
            texts
        )[:, 1]
    )

    char_scores = (
        models["char"]
        .predict_proba(
            texts
        )[:, 1]
    )

    style_scores = (
        models["style"]
        .predict_proba(
            stylometry_matrix(
                texts
            )
        )[:, 1]
    )

    component_scores = (
        np.column_stack(
            [
                word_scores,
                char_scores,
                style_scores,
            ]
        )
    )

    final_scores = (
        bundle["meta_model"]
        .predict_proba(
            meta_score_features(
                component_scores
            )
        )[:, 1]
    )

    return (
        final_scores,
        component_scores,
    )


def classify_score(
    score: float,
    human_threshold: float,
    ai_threshold: float,
) -> tuple[str, str]:
    if score >= ai_threshold:
        return (
            "Indikasi AI",
            "ai",
        )

    if score <= human_threshold:
        return (
            "Cenderung manusia",
            "human",
        )

    return (
        "Belum pasti",
        "uncertain",
    )


def weighted_average(
    values: np.ndarray,
    weights: np.ndarray,
) -> float:
    return float(
        np.average(
            values,
            weights=weights,
        )
    )


def confidence_label(
    verdict_key: str,
    word_count: int,
    model_agreement: float,
    section_consistency: float,
) -> tuple[str, int]:
    length_factor = min(
        max(
            (
                word_count
                - 40
            )
            / 360,
            0.0,
        ),
        1.0,
    )

    score = (
        0.50
        * model_agreement
        + 0.30
        * section_consistency
        + 0.20
        * length_factor
    )

    if verdict_key == "uncertain":
        score = min(
            score,
            0.54,
        )

    value = int(
        round(
            score
            * 100
        )
    )

    if value >= 76:
        return (
            "Tinggi",
            value,
        )

    if value >= 54:
        return (
            "Sedang",
            value,
        )

    return (
        "Rendah",
        value,
    )


def build_notes(
    text: str,
    word_count: int,
    chunks: list[str],
    final_scores: np.ndarray,
    component_scores: np.ndarray,
    human_threshold: float,
    ai_threshold: float,
) -> list[str]:
    notes = []

    if word_count < 80:
        notes.append(
            "Teks pendek. Hasil lebih mudah berubah jika beberapa kalimat diedit."
        )

    elif word_count < 150:
        notes.append(
            "Panjang teks cukup untuk pemeriksaan awal, tetapi dokumen yang lebih panjang biasanya lebih stabil."
        )

    language_score = (
        indonesian_signal(
            text
        )
    )

    if language_score < 0.035:
        notes.append(
            "Teks tidak terlihat dominan berbahasa Indonesia. Model ini dilatih terutama untuk Bahasa Indonesia."
        )

    component_spread = float(
        np.mean(
            np.ptp(
                component_scores,
                axis=1,
            )
        )
    )

    if component_spread > 0.36:
        notes.append(
            "Model kata, karakter, dan gaya tulis memberi hasil yang cukup berbeda."
        )

    if len(chunks) >= 2:
        ai_parts = int(
            np.sum(
                final_scores
                >= ai_threshold
            )
        )

        human_parts = int(
            np.sum(
                final_scores
                <= human_threshold
            )
        )

        if (
            ai_parts > 0
            and human_parts > 0
        ):
            notes.append(
                "Dokumen berisi bagian yang terbaca berbeda satu sama lain."
            )

    return notes


def analyze(
    text: str,
) -> dict:
    text = normalize_text(
        text
    )

    word_count = count_words(
        text
    )

    if word_count < 40:
        raise HTTPException(
            status_code=400,
            detail=(
                "Teks terlalu pendek. "
                "Masukkan minimal 40 kata "
                "agar hasil tidak terlalu mudah berubah."
            ),
        )

    bundle = get_bundle()

    metadata = bundle.get(
        "metadata",
        {},
    )

    thresholds = (
        metadata.get(
            "decision_thresholds",
            {
                "human": 0.22,
                "ai": 0.78,
            },
        )
    )

    human_threshold = float(
        thresholds.get(
            "human",
            0.22,
        )
    )

    ai_threshold = float(
        thresholds.get(
            "ai",
            0.78,
        )
    )

    chunks = chunk_text(
        text
    )

    (
        final_scores,
        component_scores,
    ) = predict_texts(
        chunks,
        bundle,
    )

    weights = np.asarray(
        [
            max(
                count_words(
                    chunk
                ),
                1,
            )
            for chunk in chunks
        ],
        dtype=np.float64,
    )

    mean_score = weighted_average(
        final_scores,
        weights,
    )

    median_score = float(
        np.median(
            final_scores
        )
    )

    if len(final_scores) == 1:
        overall_score = mean_score
    else:
        overall_score = (
            0.45
            * mean_score
            + 0.55
            * median_score
        )

    component_overall = (
        np.asarray(
            [
                weighted_average(
                    component_scores[
                        :,
                        index,
                    ],
                    weights,
                )
                for index in range(3)
            ],
            dtype=np.float64,
        )
    )

    component_spread = float(
        np.mean(
            np.ptp(
                component_scores,
                axis=1,
            )
        )
    )

    model_agreement = float(
        np.clip(
            1.0
            - component_spread
            / 0.70,
            0.0,
            1.0,
        )
    )

    if len(final_scores) > 1:
        section_std = float(
            np.std(
                final_scores
            )
        )
    else:
        section_std = 0.0

    section_consistency = float(
        np.clip(
            1.0
            - section_std
            / 0.34,
            0.0,
            1.0,
        )
    )

    ai_votes = int(
        np.sum(
            final_scores
            >= ai_threshold
        )
    )

    human_votes = int(
        np.sum(
            final_scores
            <= human_threshold
        )
    )

    uncertain_votes = (
        len(final_scores)
        - ai_votes
        - human_votes
    )

    if len(final_scores) == 1:
        (
            verdict_label,
            verdict_key,
        ) = classify_score(
            overall_score,
            human_threshold,
            ai_threshold,
        )

    else:
        minimum_votes = max(
            1,
            int(
                np.ceil(
                    len(final_scores)
                    * 0.50
                )
            ),
        )

        if (
            overall_score
            >= ai_threshold
            and ai_votes
            >= minimum_votes
        ):
            (
                verdict_label,
                verdict_key,
            ) = (
                "Indikasi AI kuat",
                "ai",
            )

        elif (
            overall_score
            <= human_threshold
            and human_votes
            >= minimum_votes
        ):
            (
                verdict_label,
                verdict_key,
            ) = (
                "Cenderung ditulis manusia",
                "human",
            )

        else:
            (
                verdict_label,
                verdict_key,
            ) = (
                "Belum cukup bukti",
                "uncertain",
            )

    (
        confidence,
        confidence_score,
    ) = confidence_label(
        verdict_key,
        word_count,
        model_agreement,
        section_consistency,
    )

    sections = []

    for index, chunk in enumerate(
        chunks
    ):
        (
            section_label,
            section_key,
        ) = classify_score(
            float(
                final_scores[
                    index
                ]
            ),
            human_threshold,
            ai_threshold,
        )

        sections.append(
            {
                "index": (
                    index + 1
                ),
                "text": chunk,
                "word_count": (
                    count_words(
                        chunk
                    )
                ),
                "score": int(
                    round(
                        float(
                            final_scores[
                                index
                            ]
                        )
                        * 100
                    )
                ),
                "label": (
                    section_label
                ),
                "label_key": (
                    section_key
                ),
                "components": {
                    "word": int(
                        round(
                            float(
                                component_scores[
                                    index,
                                    0,
                                ]
                            )
                            * 100
                        )
                    ),
                    "char": int(
                        round(
                            float(
                                component_scores[
                                    index,
                                    1,
                                ]
                            )
                            * 100
                        )
                    ),
                    "style": int(
                        round(
                            float(
                                component_scores[
                                    index,
                                    2,
                                ]
                            )
                            * 100
                        )
                    ),
                },
            }
        )

    if verdict_key == "ai":
        summary = (
            "Sebagian besar bagian melewati "
            "ambang AI dan hasil antarbagian "
            "cukup konsisten."
        )

    elif verdict_key == "human":
        summary = (
            "Sebagian besar bagian berada "
            "di bawah ambang AI yang digunakan model."
        )

    else:
        summary = (
            "Skor atau keputusan antarbagian "
            "belum cukup konsisten untuk memilih satu sisi."
        )

    notes = build_notes(
        text,
        word_count,
        chunks,
        final_scores,
        component_scores,
        human_threshold,
        ai_threshold,
    )

    evaluation = (
        metadata.get(
            "evaluation",
            {},
        )
    )

    return {
        "verdict": (
            verdict_label
        ),
        "verdict_key": (
            verdict_key
        ),
        "summary": summary,
        "ai_index": int(
            round(
                overall_score
                * 100
            )
        ),
        "confidence": (
            confidence
        ),
        "confidence_score": (
            confidence_score
        ),
        "model_agreement": int(
            round(
                model_agreement
                * 100
            )
        ),
        "section_consistency": int(
            round(
                section_consistency
                * 100
            )
        ),
        "word_count": (
            word_count
        ),
        "character_count": len(
            text
        ),
        "section_count": len(
            chunks
        ),
        "section_summary": {
            "ai": ai_votes,
            "human": (
                human_votes
            ),
            "uncertain": (
                uncertain_votes
            ),
        },
        "components": {
            "word": int(
                round(
                    float(
                        component_overall[
                            0
                        ]
                    )
                    * 100
                )
            ),
            "char": int(
                round(
                    float(
                        component_overall[
                            1
                        ]
                    )
                    * 100
                )
            ),
            "style": int(
                round(
                    float(
                        component_overall[
                            2
                        ]
                    )
                    * 100
                )
            ),
        },
        "style": (
            style_snapshot(
                text
            )
        ),
        "notes": notes,
        "sections": sections,
        "model": {
            "name": (
                metadata.get(
                    "model_name",
                    "DeteksiAI Hybrid ID",
                )
            ),
            "version": (
                metadata.get(
                    "model_version",
                    "2.0",
                )
            ),
            "sources": (
                metadata.get(
                    "sources",
                    {},
                )
            ),
            "test_samples": (
                metadata
                .get(
                    "samples",
                    {},
                )
                .get(
                    "test"
                )
            ),
            "selective_accuracy": (
                evaluation.get(
                    "selective_accuracy"
                )
            ),
            "selective_coverage": (
                evaluation.get(
                    "selective_coverage"
                )
            ),
            "ai_false_positive_rate": (
                evaluation.get(
                    "ai_false_positive_rate"
                )
            ),
            "trained_at": (
                metadata.get(
                    "trained_at"
                )
            ),
        },
        "thresholds": {
            "human": int(
                round(
                    human_threshold
                    * 100
                )
            ),
            "ai": int(
                round(
                    ai_threshold
                    * 100
                )
            ),
        },
    }


@app.get(
    "/",
    response_class=HTMLResponse,
)
def home(
    request: Request,
):
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={},
    )


@app.get(
    "/favicon.ico",
    include_in_schema=False,
)
def favicon():
    return Response(
        status_code=204
    )


@app.get(
    "/api/health"
)
def health():
    if not MODEL_PATH.exists():
        return {
            "status": "ok",
            "model_ready": False,
            "model_version": None,
        }

    try:
        bundle = get_bundle()

        version = (
            bundle
            .get(
                "metadata",
                {},
            )
            .get(
                "model_version"
            )
        )

        return {
            "status": "ok",
            "model_ready": True,
            "model_version": version,
        }

    except HTTPException as exception:
        return {
            "status": "ok",
            "model_ready": False,
            "model_version": None,
            "detail": (
                exception.detail
            ),
        }


@app.post(
    "/api/analyze"
)
async def analyze_endpoint(
    text: Optional[str] = Form(
        default=None
    ),
    file: Optional[UploadFile] = File(
        default=None
    ),
):
    content = text or ""

    if (
        file
        and file.filename
    ):
        raw = await file.read(
            MAX_FILE_BYTES + 1
        )

        if (
            len(raw)
            > MAX_FILE_BYTES
        ):
            raise HTTPException(
                status_code=413,
                detail=(
                    "Ukuran file maksimal 8 MB."
                ),
            )

        extracted = extract_text(
            file.filename,
            raw,
        )

        if not extracted.strip():
            raise HTTPException(
                status_code=400,
                detail=(
                    "Tidak ada teks yang dapat dibaca "
                    "dari dokumen. PDF hasil scan gambar "
                    "belum didukung."
                ),
            )

        if content.strip():
            content = (
                f"{content}\n\n"
                f"{extracted}"
            )
        else:
            content = extracted

    if not content.strip():
        raise HTTPException(
            status_code=400,
            detail=(
                "Masukkan teks atau pilih dokumen terlebih dahulu."
            ),
        )

    return analyze(
        content
    )