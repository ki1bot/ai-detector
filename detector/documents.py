import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from detector.config import ALLOWED_EXTENSIONS


class DocumentError(ValueError):
    pass


def extract_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()

        if value:
            blocks.append(value)

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                blocks.append(" | ".join(values))

    return "\n\n".join(blocks)


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []

    for page in reader.pages:
        value = page.extract_text() or ""
        value = re.sub(
            r"[ \t]+",
            " ",
            value,
        ).strip()

        if value:
            pages.append(value)

    return "\n\n".join(pages)


def extract_text(
    filename: str,
    content: bytes,
) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix not in ALLOWED_EXTENSIONS:
        raise DocumentError(
            "Format file harus TXT, MD, DOCX, atau PDF."
        )

    try:
        if suffix in {".txt", ".md"}:
            return content.decode(
                "utf-8",
                errors="replace",
            )

        if suffix == ".docx":
            return extract_docx(content)

        return extract_pdf(content)

    except DocumentError:
        raise

    except Exception as exception:
        raise DocumentError(
            f"Dokumen tidak dapat dibaca: {exception}"
        ) from exception